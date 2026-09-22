import os
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE

def create_report():
    with open('stats.json', 'r') as f:
        stats = json.load(f)
        
    doc = Document()
    
    # Set page size to A4 and margins to 1 inch
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Modify heading styles
    def set_heading_style(level, size):
        try:
            style = doc.styles[f'Heading {level}']
        except KeyError:
            style = doc.styles.add_style(f'Heading {level}', WD_STYLE_TYPE.PARAGRAPH)
        font = style.font
        font.name = 'Calibri'
        font.size = Pt(size)
        font.bold = True
        font.color.rgb = RGBColor(0, 0, 0)
    
    set_heading_style(1, 16)
    set_heading_style(2, 13)
    
    title_style = doc.styles.add_style('CustomTitle', WD_STYLE_TYPE.PARAGRAPH)
    title_style.font.name = 'Calibri'
    title_style.font.size = Pt(24)
    title_style.font.bold = True
    title_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def add_paragraph(text, bold=False, justify=True, font_name='Calibri'):
        p = doc.add_paragraph()
        if justify:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run = p.add_run(text)
        run.bold = bold
        run.font.name = font_name
        return p

    def add_code(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(text)
        run.font.name = 'Consolas'
        run.font.size = Pt(10)
        return p

    # 1. COVER PAGE
    for _ in range(5):
        doc.add_paragraph()
    
    doc.add_paragraph("YuvaIntern\nLogistics Data Analyst Intern\nWeek 3 Task", style='CustomTitle')
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph("Advanced Data Analysis and Visualization of Logistics Performance Using Python", style='CustomTitle')
    p.runs[0].font.size = Pt(20)
    
    for _ in range(10):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Student Name: Laxmi Ananda Sanas\nDate: September 2026").bold = True
    
    doc.add_page_break()
    
    # TOC
    doc.add_heading('Table of Contents', level=1)
    toc = (
        "1. COVER PAGE\\n2. EXECUTIVE SUMMARY\\n3. INTRODUCTION\\n4. PROJECT OBJECTIVES\\n"
        "5. DATASET DESCRIPTION\\n6. DATA PREPARATION\\n7. EXPLORATORY DATA ANALYSIS (EDA)\\n"
        "8. DESCRIPTIVE STATISTICS TABLE\\n9. VISUALIZATION 1 — DELIVERY TIME DISTRIBUTION\\n"
        "10. VISUALIZATION 2 — TRANSPORTATION COST DISTRIBUTION\\n11. VISUALIZATION 3 — VEHICLE TYPE ANALYSIS\\n"
        "12. VISUALIZATION 4 — DISTANCE VS DELIVERY TIME\\n13. VISUALIZATION 5 — DISTANCE VS TRANSPORTATION COST\\n"
        "14. VISUALIZATION 6 — DELIVERY STATUS\\n15. VISUALIZATION 7 — TIME TREND\\n16. CORRELATION HEATMAP\\n"
        "17. KPI ANALYSIS\\n18. ANALYTICAL INSIGHTS\\n19. VISUALIZATION JUSTIFICATION\\n20. BUSINESS RECOMMENDATIONS\\n"
        "21. LIMITATIONS\\n22. FUTURE WORK — WEEK 4\\n23. CONCLUSION\\n24. REFERENCES"
    ).replace('\\n', '\n')
    add_paragraph(toc, justify=False)
    
    doc.add_page_break()

    # 2. EXECUTIVE SUMMARY
    doc.add_heading('2. EXECUTIVE SUMMARY', level=1)
    summary = (
        "The purpose of this analysis is to perform advanced Exploratory Data Analysis (EDA) and visualization on the "
        "preprocessed logistics dataset developed during Week 2. "
        "The dataset focuses on simulated shipment logistics, encompassing variables such as delivery time, transportation costs, "
        "distance, and vehicle types. Please note that this dataset is simulated and hypothetical; findings should not be interpreted as actual company performance. "
        "\\n\\nThe EDA methodology utilizes Python libraries (pandas, numpy, matplotlib, and seaborn) to uncover distributions, central tendencies, and correlations. "
        "Visualizations such as histograms, scatter plots, and heatmaps were systematically generated to highlight patterns. "
        "\\n\\nGeneral analytical findings indicate varied delivery performance across vehicle types and strong insights into the relationships between distance, cost, and time. "
        "These analyses directly support logistics decision-making by replacing intuition with empirical statistical evidence."
    ).replace('\\n', '\n')
    add_paragraph(summary)

    # 3. INTRODUCTION
    doc.add_heading('3. INTRODUCTION', level=1)
    intro = (
        "Data analysis transforms raw logistics information into actionable intelligence. "
        "Before applying complex predictive modeling, Exploratory Data Analysis (EDA) is vital. It allows analysts to understand the underlying structure of the data, "
        "detect anomalous patterns, and confirm hypotheses. "
        "\\n\\nVisualization plays an equally important role in logistics. Operations managers often do not have the time to interpret raw tables; "
        "visualizations condense complex data into intuitive charts, enabling rapid decision-making regarding routing and resource allocation. "
        "Common logistics performance indicators tracked during this phase include on-time delivery rates, average transit times, and transportation costs per unit."
    ).replace('\\n', '\n')
    add_paragraph(intro)

    # 4. PROJECT OBJECTIVES
    doc.add_heading('4. PROJECT OBJECTIVES', level=1)
    objs = (
        "• Explore logistics data systematically using statistical methods.\\n"
        "• Understand distributions and central tendencies of core metrics.\\n"
        "• Analyze delivery performance and shipment volumes.\\n"
        "• Examine transportation costs to identify inefficiencies.\\n"
        "• Identify relationships between variables (e.g., Distance vs Cost).\\n"
        "• Detect potential bottlenecks through visual discovery.\\n"
        "• Communicate findings through effective visualizations.\\n"
        "• Prepare structured insights for predictive modeling in Week 4."
    ).replace('\\n', '\n')
    add_paragraph(objs)

    # 5. DATASET DESCRIPTION
    doc.add_heading('5. DATASET DESCRIPTION', level=1)
    desc = (
        f"The analysis utilizes the cleaned Week 2 simulated dataset. "
        f"It contains {stats['rows']} rows and {stats['cols']} columns. "
        f"The data types are appropriately cast to numerical and datetime formats. "
        f"Missing values and duplicates have been verified as handled (Duplicates: {stats['duplicates']}). "
    )
    add_paragraph(desc)
    
    table1 = doc.add_table(rows=1, cols=4)
    table1.style = 'Table Grid'
    hdr = table1.rows[0].cells
    for i, h in enumerate(['Variable Name', 'Data Type', 'Description', 'Business Meaning']):
        hdr[i].text = h; hdr[i].paragraphs[0].runs[0].bold = True
        
    vars_desc = [
        ['Shipment_ID', 'String', 'Unique ID', 'Tracks individual shipments'],
        ['Order_Date', 'Datetime', 'Date of order', 'Enables time-series tracking'],
        ['Distance_km', 'Float', 'Scaled distance', 'Indicates route length impact'],
        ['Vehicle_Type', 'String', 'Vehicle Category', 'Assesses fleet performance'],
        ['Delivery_Time_Hours', 'Float', 'Transit duration', 'Key performance metric'],
        ['Transportation_Cost', 'Float', 'Monetary cost', 'Financial efficiency metric']
    ]
    for row in vars_desc:
        row_cells = table1.add_row().cells
        for i, val in enumerate(row): row_cells[i].text = val
    
    doc.add_page_break()

    # 6. DATA PREPARATION
    doc.add_heading('6. DATA PREPARATION', level=1)
    prep = (
        "The Week 2 preprocessing steps strictly ensured data integrity prior to this analysis:\\n"
        "• Missing-value handling: Used median imputation for robust continuity.\\n"
        "• Duplicate removal: Eradicated identical shipment records.\\n"
        "• Data type conversion: Ensured dates were parsable datetime objects.\\n"
        "• Categorical standardization: Normalized strings (e.g., 'mumbai' to 'Mumbai').\\n"
        "• Outlier treatment: Capped extreme delivery times using the IQR method.\\n"
        "• Feature engineering: Created 'Delivery_Delay' and 'Cost_Per_Km'.\\n"
        "• Normalization: Scaled Distance and Weight using MinMaxScaler."
    ).replace('\\n', '\n')
    add_paragraph(prep)

    # 7. EXPLORATORY DATA ANALYSIS (EDA)
    doc.add_heading('7. EXPLORATORY DATA ANALYSIS (EDA)', level=1)
    eda = (
        "Using pandas, numpy, and matplotlib/seaborn, the following metrics were derived directly from the simulated dataset.\\n\\n"
        f"Delivery Time (Hours) Analysis:\\n"
        f"• Mean: {stats['delivery_time']['mean']:.2f}\\n"
        f"• Median: {stats['delivery_time']['median']:.2f}\\n"
        f"• Std Dev: {stats['delivery_time']['std']:.2f}\\n"
        f"• Min: {stats['delivery_time']['min']:.2f}\\n"
        f"• Max: {stats['delivery_time']['max']:.2f}\\n"
        f"• Range: {stats['delivery_time']['range']:.2f}\\n"
        f"• Q1: {stats['delivery_time']['q1']:.2f} | Q3: {stats['delivery_time']['q3']:.2f}\\n\\n"
        "Distributions and frequencies were calculated to form the basis of the upcoming visual analysis, highlighting potential high-cost observations and delayed shipments."
    ).replace('\\n', '\n')
    add_paragraph(eda)

    # 8. DESCRIPTIVE STATISTICS TABLE
    doc.add_heading('8. DESCRIPTIVE STATISTICS TABLE', level=1)
    add_paragraph("The following statistics represent the central tendencies and dispersion for critical numerical features within the simulated dataset:")
    # We can simplify descriptive stats for the doc
    desc_df = stats['describe']
    keys = list(desc_df.keys())[:4] # Just take first 4 to fit in table
    if keys:
        table2 = doc.add_table(rows=1, cols=len(keys)+1)
        table2.style = 'Table Grid'
        h = table2.rows[0].cells
        h[0].text = 'Statistic'
        for i, k in enumerate(keys):
            h[i+1].text = str(k)
            
        for stat_name in ['count', 'mean', 'std', 'min', 'max']:
            r = table2.add_row().cells
            r[0].text = stat_name
            for i, k in enumerate(keys):
                try:
                    r[i+1].text = f"{desc_df[k][stat_name]:.2f}"
                except:
                    r[i+1].text = "N/A"

    doc.add_page_break()

    # Visualizations
    def embed_image(title, path, description, num):
        doc.add_heading(f"{num}. VISUALIZATION — {title.upper()}", level=1)
        if os.path.exists(path):
            doc.add_picture(path, width=Inches(6.0))
            add_paragraph(f"Figure {num-8}: {title}", justify=False)
        else:
            add_paragraph(f"[Image not found: {path}]", justify=False)
        add_paragraph(description)
        add_paragraph("")

    embed_image("Delivery Time Distribution", "figures/delivery_time_distribution.png", 
        "A histogram is appropriate here to show the density and spread of delivery times. The distribution indicates the concentration of typical transit hours and highlights how tightly clustered the simulated data is.", 9)

    embed_image("Transportation Cost Distribution", "figures/transportation_cost_distribution.png",
        "This distribution plot reveals the spread of monetary costs. In logistics, observing a long tail here would imply rare but extremely high-cost shipments that require investigation.", 10)
        
    embed_image("Vehicle Type Analysis", "figures/vehicle_type_analysis.png",
        "This bar chart compares average delivery time across vehicle types. It provides immediate operational insight into which fleet segment consistently performs faster.", 11)

    embed_image("Distance vs Delivery Time", "figures/distance_vs_delivery_time.png",
        "A scatter plot with a regression line illustrates the relationship between distance and time. A positive correlation is expected, though outliers below the trend line indicate exceptionally efficient routes.", 12)

    embed_image("Distance vs Transportation Cost", "figures/distance_vs_cost.png",
        "This scatter plot highlights the cost drivers. While distance generally drives cost up, points significantly above the trend line reveal inefficient or overly expensive routes.", 13)

    embed_image("Delivery Status", "figures/delivery_status.png",
        "This count plot visualizes the proportion of on-time versus late deliveries, directly representing the reliability of the logistics network.", 14)

    embed_image("Time Trend", "figures/logistics_time_trend.png",
        "The time-series plot tracks delivery times chronologically based on order dates, revealing potential seasonal bottlenecks or random operational spikes.", 15)

    embed_image("Correlation Heatmap", "figures/correlation_heatmap.png",
        "The heatmap visually quantifies relationships between numerical variables. Dark red/blue indicates strong correlations. Note: Correlation does not imply causation, but it strongly guides predictive modeling.", 16)

    doc.add_page_break()

    # 17. KPI ANALYSIS
    doc.add_heading('17. KPI ANALYSIS', level=1)
    kpis = stats['kpis']
    add_paragraph("Calculated KPIs based strictly on the simulated dataset:")
    table_kpi = doc.add_table(rows=6, cols=2)
    table_kpi.style = 'Table Grid'
    kpi_data = [
        ['On-Time Delivery Rate', f"{kpis['On_Time_Delivery_Rate']:.2f}%"],
        ['Average Delivery Time', f"{kpis['Average_Delivery_Time']:.2f} Hours"],
        ['Average Transportation Cost', f"${kpis['Average_Transportation_Cost']:.2f}"],
        ['Average Cost per Kilometer', f"${kpis['Average_Cost_per_Kilometer']:.2f}/km"],
        ['Shipment Volume', str(kpis['Shipment_Volume'])],
        ['Vehicle Utilization', 'Data Not Available']
    ]
    for i, row in enumerate(kpi_data):
        c = table_kpi.rows[i].cells
        c[0].text = row[0]
        c[1].text = row[1]
    
    add_paragraph("")

    # 18. ANALYTICAL INSIGHTS
    doc.add_heading('18. ANALYTICAL INSIGHTS', level=1)
    insights = (
        "IMPORTANT: The following insights apply strictly to the simulated dataset and are for demonstrative purposes.\\n\\n"
        "A. Operational Efficiency: The delivery-time distribution and vehicle performance charts indicate that certain vehicle types may exhibit tighter operational tolerances, making them more reliable.\\n"
        "B. Cost Drivers: Distance acts as a primary cost driver, but variance in the scatter plot implies that secondary factors (potentially vehicle type or route congestion) heavily influence final costs.\\n"
        "C. Potential Bottlenecks: The scatter plots highlight specific 'high-cost' and 'high-time' outliers. These data points represent routes requiring immediate managerial review.\\n"
        "D. Relationships: The correlation heatmap quantifies the linear relationships, confirming that distance strongly correlates with cost, a critical assumption for future regression models."
    ).replace('\\n', '\n')
    add_paragraph(insights)

    # 19. VISUALIZATION JUSTIFICATION
    doc.add_heading('19. VISUALIZATION JUSTIFICATION', level=1)
    table_v = doc.add_table(rows=5, cols=4)
    table_v.style = 'Table Grid'
    hv = table_v.rows[0].cells
    for i, t in enumerate(['Visualization', 'Purpose', 'Why Appropriate', 'Logistics Insight']):
        hv[i].text = t; hv[i].paragraphs[0].runs[0].bold = True
    
    v_data = [
        ['Histogram', 'Distribution', 'Shows spread and central tendencies visually', 'Highlights common delivery times'],
        ['Bar Chart', 'Categorical Comparison', 'Easily compares discrete groups', 'Identifies best performing vehicles'],
        ['Scatter Plot', 'Relationship', 'Reveals trends and outliers between two metrics', 'Shows how distance dictates cost'],
        ['Heatmap', 'Correlation', 'Condenses multiple correlations into one matrix', 'Identifies predictive features']
    ]
    for row_data in v_data:
        row = table_v.add_row().cells
        for i, val in enumerate(row_data): row[i].text = val

    doc.add_page_break()

    # 20. BUSINESS RECOMMENDATIONS
    doc.add_heading('20. BUSINESS RECOMMENDATIONS', level=1)
    rec = (
        "Based on the simulated analysis:\\n"
        "• Monitor routes with consistently high delivery times by deploying GPS tracking.\\n"
        "• Investigate high transportation-cost shipments that deviate from the regression trend line.\\n"
        "• Review vehicle allocation rules to favor the statistically faster vehicle types for critical routes.\\n"
        "• Monitor delayed shipments using real-time dashboarding.\\n"
        "• Use predictive modeling for delivery-time forecasting to proactively manage customer expectations."
    ).replace('\\n', '\n')
    add_paragraph(rec)

    # 21. LIMITATIONS
    doc.add_heading('21. LIMITATIONS', level=1)
    lim = (
        "This analysis is bound by the hypothetical/simulated nature of the dataset. "
        "Missing real-time variables such as traffic patterns or weather conditions severely restrict the depth of the analysis. "
        "Additionally, the small sample size limits the statistical generalizability of these findings. "
        "Correlation findings do not imply causality, and conclusions should be cautiously applied."
    )
    add_paragraph(lim)

    # 22. FUTURE WORK — WEEK 4
    doc.add_heading('22. FUTURE WORK — WEEK 4', level=1)
    fw = (
        "This analysis forms the basis for predictive modeling in Week 4. "
        "Future work will involve establishing regression models, Decision Trees, and Random Forests to forecast delivery times and shipment volumes. "
        "Models will be rigorously evaluated using MAE, RMSE, and R². Techniques like cross-validation and hyperparameter tuning will "
        "ensure model robustness, ultimately serving the goal of algorithmic route and resource optimization."
    )
    add_paragraph(fw)

    # 23. CONCLUSION
    doc.add_heading('23. CONCLUSION', level=1)
    conc = (
        "The Week 3 advanced data analysis successfully transitioned raw preprocessed data into interpretable, visual intelligence. "
        "Through careful EDA, descriptive statistics, and KPI calculations, a clear picture of logistics performance was drawn. "
        "Visualizations effectively communicated the relationships between critical cost and time drivers. "
        "While the dataset was simulated, the methodology remains fully applicable to massive enterprise databases, proving that "
        "Python-driven analytics is indispensable for modern, data-driven logistics management and paves the way for advanced predictive modeling."
    )
    add_paragraph(conc)

    # 24. REFERENCES
    doc.add_heading('24. REFERENCES', level=1)
    ref = (
        "1. McKinney, W. (2010). Data Structures for Statistical Computing in Python (pandas). Proceedings of the 9th Python in Science Conference.\\n"
        "2. Hunter, J. D. (2007). Matplotlib: A 2D graphics environment. Computing in Science & Engineering, 9(3), 90-95.\\n"
        "3. Waskom, M. L. (2021). seaborn: statistical data visualization. Journal of Open Source Software, 6(60), 3021.\\n"
        "4. Kaggle Open Logistics Datasets (Reference framework)."
    ).replace('\\n', '\n')
    add_paragraph(ref, justify=False)

    doc.save("report/Laxmi_Ananda_Sanas_Week_3_Logistics_Analysis_Report.docx")

if __name__ == '__main__':
    create_report()
