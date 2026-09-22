import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.style import WD_STYLE_TYPE

def create_report():
    doc = Document()
    
    # Set page size to A4 and margins to 1 inch
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    
    # Modify heading styles
    def set_heading_style(level, size):
        try:
            style = doc.styles[f'Heading {level}']
        except KeyError:
            style = doc.styles.add_style(f'Heading {level}', WD_STYLE_TYPE.PARAGRAPH)
        font = style.font
        font.name = 'Calibri'
        font.size = Pt(size)
        font.bold = True
        font.color.rgb = RGBColor(0, 0, 0)
    
    set_heading_style(1, 16)
    set_heading_style(2, 13)
    
    # Custom Title style
    title_style = doc.styles.add_style('CustomTitle', WD_STYLE_TYPE.PARAGRAPH)
    title_style.font.name = 'Calibri'
    title_style.font.size = Pt(24)
    title_style.font.bold = True
    title_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def add_paragraph(text, bold=False, justify=True, font_name='Calibri'):
        p = doc.add_paragraph()
        if justify:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run = p.add_run(text)
        run.bold = bold
        run.font.name = font_name
        return p

    def add_code(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(text)
        run.font.name = 'Consolas'
        run.font.size = Pt(10)
        return p

    # 1. COVER PAGE
    for _ in range(5):
        doc.add_paragraph()
    
    p = doc.add_paragraph("YuvaIntern\nLogistics Data Analyst Intern\nWeek 1 Task", style='CustomTitle')
    
    for _ in range(3):
        doc.add_paragraph()
        
    p = doc.add_paragraph("Logistics Delivery Performance and Route Optimization Using Data Science and Python", style='CustomTitle')
    p.runs[0].font.size = Pt(20)
    
    for _ in range(10):
        doc.add_paragraph()
        
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Student Name: Laxmi Ananda Sanas\nDate: September 2026").bold = True
    
    doc.add_page_break()
    
    # TOC placeholder
    doc.add_heading('Table of Contents', level=1)
    add_paragraph("1. COVER PAGE\n2. EXECUTIVE SUMMARY\n3. INTRODUCTION\n4. BACKGROUND RESEARCH\n5. PROJECT SCENARIO\n6. PROBLEM STATEMENT\n7. PROJECT OBJECTIVES\n8. KEY PERFORMANCE INDICATORS (KPIs)\n9. DATA REQUIREMENTS\n10. PUBLIC DATASET / DATA RESEARCH\n11. DATA SCIENCE METHODOLOGIES\n12. STRATEGIC ANALYTICAL ROADMAP\n13. DATA CLEANING PLAN\n14. EXPLORATORY DATA ANALYSIS PLAN\n15. PYTHON CODE ILLUSTRATIONS\n16. KPI CALCULATION EXAMPLE\n17. EXPECTED ANALYTICAL INSIGHTS\n18. EXPECTED OUTCOMES\n19. FUTURE EXTENSION\n20. LIMITATIONS\n21. CONCLUSION\n22. REFERENCES", justify=False)
    
    doc.add_page_break()

    # 2. EXECUTIVE SUMMARY
    doc.add_heading('2. EXECUTIVE SUMMARY', level=1)
    summary = (
        "The purpose of this project is to develop a strategic plan for analyzing logistics delivery performance and route optimization. "
        "The logistics industry faces continuous challenges such as delivery delays, high transportation costs, and inefficient routes. "
        "Addressing these challenges requires a robust analytical approach. Data analytics is exceptionally useful in logistics because it "
        "enables companies to shift from reactive operations to proactive, data-driven decision-making. By leveraging historical shipment data, "
        "organizations can uncover hidden patterns, optimize vehicle utilization, and significantly improve on-time delivery rates. "
        "\\n\\nTo effectively measure and monitor operational performance, key performance indicators (KPIs) such as On-Time Delivery Rate, Average Delivery Time, "
        "Transportation Cost per Shipment, Shipment Volume, and Vehicle Utilization Rate have been selected. These metrics provide a comprehensive view of logistics efficiency. "
        "\\n\\nThe proposed approach involves utilizing Python and its data science ecosystem (including pandas, numpy, and seaborn) to perform exploratory data analysis, "
        "correlation mapping, and eventually predictive modeling. Through a systematic methodology encompassing data cleaning, feature engineering, and statistical analysis, "
        "the project aims to build a foundation for advanced route optimization models. "
        "\\n\\nThe expected business benefits of this strategic roadmap include enhanced visibility into the supply chain, identification of costly bottlenecks, "
        "improved resource allocation, and a tangible reduction in transportation costs. Ultimately, this data-driven approach is expected to improve delivery reliability "
        "and foster a more resilient logistics operation."
    ).replace('\\n', '\n')
    add_paragraph(summary)
    
    # 3. INTRODUCTION
    doc.add_heading('3. INTRODUCTION', level=1)
    intro = (
        "Logistics and supply chain operations involve the complex coordination of sourcing, storing, and transporting goods from origin to destination. "
        "It encompasses inventory management, fleet operations, warehousing, and final-mile delivery. "
        "\\n\\nDelivery efficiency is paramount in today's fast-paced business environment. Customers expect rapid, reliable shipments, and failure to meet these expectations "
        "can lead to lost business and damaged brand reputation. Furthermore, efficient delivery directly impacts the bottom line by minimizing fuel consumption and operational overhead. "
        "\\n\\nHowever, logistics providers face numerous common challenges. These include unexpected delays due to traffic or weather, high transportation costs, inefficient delivery routes, "
        "poor resource allocation leading to underutilized vehicles, and inconsistent seasonal demand. "
        "\\n\\nData science provides a powerful toolkit to support logistics decision-making and address these challenges. By analyzing large volumes of transportation data, "
        "companies can predict potential delays, optimize routing schedules algorithmically, and dynamically allocate resources based on real-time insights, thereby transforming operations."
    ).replace('\\n', '\n')
    add_paragraph(intro)
    
    # 4. BACKGROUND RESEARCH
    doc.add_heading('4. BACKGROUND RESEARCH', level=1)
    research = (
        "Logistics analytics has become a cornerstone of modern supply chain management. According to research by institutions like the World Bank, efficient logistics infrastructure is highly correlated with national economic competitiveness. "
        "Supply chain data—encompassing shipment records, transit times, and vehicle telemetry—provides the raw material for this analytical transformation. "
        "\\n\\nTransportation and delivery performance are critical metrics extensively studied in operations research. Optimization of these areas relies heavily on accurate historical data to forecast future trends. "
        "Route optimization, a classic application of the Traveling Salesperson Problem (TSP) and Vehicle Routing Problem (VRP), utilizes mathematical and heuristic approaches to minimize travel distance and time. "
        "\\n\\nFurthermore, inventory and resource planning benefit immensely from predictive analytics, allowing companies to position assets effectively ahead of demand spikes. "
        "The importance of data-driven decision-making in logistics cannot be overstated; it reduces reliance on intuition, lowers costs, and provides a distinct competitive advantage in a volatile market."
    ).replace('\\n', '\n')
    add_paragraph(research)

    # 5. PROJECT SCENARIO
    doc.add_heading('5. PROJECT SCENARIO', level=1)
    scenario = (
        "(Note: This is a hypothetical project scenario created for academic and internship purposes.)\\n\\n"
        "SwiftLogistics is a hypothetical mid-sized logistics company that handles a large volume of daily shipments across multiple major cities. "
        "The shipment process involves picking up goods from central warehouses and distributing them to regional hubs and end customers. "
        "\\n\\nThe company operates delivery routes that span both congested urban environments and long-haul intercity highways. "
        "Its fleet consists of a mix of light commercial vehicles for last-mile delivery and heavy-duty trucks for intercity transport. "
        "Customers range from individual consumers expecting next-day ecommerce deliveries to corporate clients requiring bulk shipments. "
        "\\n\\nThe daily shipment volume averages roughly 10,000 packages. However, SwiftLogistics is experiencing rising transportation costs and inconsistent delivery times. "
        "Possible operational problems include traffic bottlenecks on specific routes, suboptimal loading leading to low vehicle utilization, and a lack of real-time visibility into factors causing late deliveries."
    ).replace('\\n', '\n')
    add_paragraph(scenario)

    # 6. PROBLEM STATEMENT
    doc.add_heading('6. PROBLEM STATEMENT', level=1)
    prob = (
        "SwiftLogistics is facing challenges with delivery reliability and escalating operational expenses. "
        "The main problem is a lack of data-driven insight into route efficiency and the variables influencing shipment delays. "
        "\\n\\nSpecifically, the company wants to:\\n"
        "• Reduce late deliveries to improve customer satisfaction.\\n"
        "• Understand factors affecting delivery time, such as distance, traffic, and weather.\\n"
        "• Identify high-cost routes that drain profitability.\\n"
        "• Improve vehicle and resource utilization to ensure fleet assets are maximized.\\n"
        "• Support better route planning through analytical insights.\\n"
        "• Improve overall logistics efficiency to maintain a competitive edge."
    ).replace('\\n', '\n')
    add_paragraph(prob)
    
    doc.add_page_break()

    # 7. PROJECT OBJECTIVES
    doc.add_heading('7. PROJECT OBJECTIVES', level=1)
    obj = (
        "1. Analyze logistics delivery performance using historical shipment data.\\n"
        "2. Identify important factors affecting delivery time through statistical analysis.\\n"
        "3. Monitor key logistics KPIs to establish baseline operational performance.\\n"
        "4. Identify inefficient and high-cost routes requiring optimization.\\n"
        "5. Apply Python-based data analysis techniques to process and visualize logistics datasets.\\n"
        "6. Develop a strategic roadmap for future predictive modeling and algorithmic route optimization."
    ).replace('\\n', '\n')
    add_paragraph(obj)

    # 8. KEY PERFORMANCE INDICATORS (KPIs)
    doc.add_heading('8. KEY PERFORMANCE INDICATORS (KPIs)', level=1)
    
    table = doc.add_table(rows=6, cols=5)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    headers = ['KPI', 'Definition', 'Formula', 'Why it Matters', 'Decision-Making Support']
    for i, header in enumerate(headers):
        hdr_cells[i].text = header
        hdr_cells[i].paragraphs[0].runs[0].bold = True
        
    kpis = [
        ['On-Time Delivery Rate', 'Percentage of shipments delivered on or before the expected time.', '(On-Time Deliveries / Total Deliveries) x 100', 'Measures reliability and customer satisfaction.', 'Helps identify regions or routes needing intervention.'],
        ['Average Delivery Time', 'The mean time taken from shipment origin to destination.', 'Total Delivery Time / Total Shipments', 'Indicates the overall speed of the logistics network.', 'Sets benchmarks for operational efficiency.'],
        ['Transportation Cost per Shipment', 'The average cost incurred to transport a single shipment.', 'Total Transportation Cost / Total Shipments', 'Directly impacts profit margins.', 'Highlights expensive routes for cost-reduction strategies.'],
        ['Shipment Volume', 'Total number of items or orders processed in a specific period.', 'Count of Total Shipments', 'Measures operational scale and demand.', 'Assists in capacity planning and resource allocation.'],
        ['Vehicle Utilization Rate', "Percentage of a vehicle's capacity used during transit.", '(Actual Load Volume / Total Vehicle Capacity) x 100', 'Indicates if fleet assets are being wasted.', 'Guides consolidation of shipments to reduce trips.']
    ]
    
    for row_data in kpis:
        row_cells = table.add_row().cells
        for i, val in enumerate(row_data):
            row_cells[i].text = val
            
    add_paragraph("") # Spacing

    # 9. DATA REQUIREMENTS
    doc.add_heading('9. DATA REQUIREMENTS', level=1)
    add_paragraph("To achieve the project objectives, a comprehensive logistics dataset is required. The table below outlines the proposed variables:")
    
    var_table = doc.add_table(rows=1, cols=2)
    var_table.style = 'Table Grid'
    hdr_cells = var_table.rows[0].cells
    hdr_cells[0].text = 'Variable'
    hdr_cells[1].text = 'Purpose'
    hdr_cells[0].paragraphs[0].runs[0].bold = True
    hdr_cells[1].paragraphs[0].runs[0].bold = True
    
    variables = [
        ['Shipment_ID', 'Unique identifier for each shipment record.'],
        ['Order_Date', 'The date the shipment was initiated, used for temporal analysis.'],
        ['Origin', 'Starting location of the shipment.'],
        ['Destination', 'Ending location of the shipment, used for route mapping.'],
        ['Distance_km', 'Physical distance between origin and destination.'],
        ['Vehicle_Type', 'Categorization of the delivery vehicle (e.g., Truck, Van).'],
        ['Shipment_Weight_kg', 'Weight of the package, used for capacity planning.'],
        ['Shipment_Volume', 'Physical volume of the package.'],
        ['Delivery_Time_Hours', 'Actual time taken to deliver the shipment.'],
        ['Expected_Delivery_Hours', 'The promised delivery timeframe.'],
        ['Transportation_Cost', 'Total monetary cost incurred for the shipment.'],
        ['Fuel_Cost', 'Estimated fuel expense for the route.'],
        ['Delivery_Status', 'Categorical status (e.g., On-Time, Late).'],
        ['Traffic_Level', 'Categorical condition of traffic (e.g., Low, High).'],
        ['Weather_Condition', 'Environmental conditions during transit (e.g., Clear, Rain).']
    ]
    
    for row_data in variables:
        row_cells = var_table.add_row().cells
        row_cells[0].text = row_data[0]
        row_cells[1].text = row_data[1]

    doc.add_page_break()

    # 10. PUBLIC DATASET / DATA RESEARCH
    doc.add_heading('10. PUBLIC DATASET / DATA RESEARCH', level=1)
    data_research = (
        "For an authentic analysis, this project can leverage publicly available transportation and supply chain datasets. "
        "Sources such as the Kaggle data science community provide numerous simulated supply chain logs and retail delivery datasets. "
        "Additionally, the UCI Machine Learning Repository contains datasets related to logistics performance and daily demand forecasting. "
        "Government open-data portals and World Bank logistics indices also provide macroeconomic context regarding transportation efficiency. "
        "\\n\\nThese platforms offer structured tabular data containing shipment IDs, transit times, distances, and delays. "
        "This data is highly relevant as it closely mimics the hypothetical scenario of SwiftLogistics. By utilizing such datasets, "
        "we can perform realistic exploratory data analysis, compute the defined KPIs, and lay the groundwork for machine learning models."
    ).replace('\\n', '\n')
    add_paragraph(data_research)

    # 11. DATA SCIENCE METHODOLOGIES
    doc.add_heading('11. DATA SCIENCE METHODOLOGIES', level=1)
    methodologies = (
        "A. Exploratory Data Analysis (EDA): This involves visualizing and summarizing the main characteristics of the dataset. "
        "In logistics, EDA helps identify trends like peak shipping days or distribution of late deliveries, providing a foundational understanding of operations.\\n\\n"
        "B. Regression: A statistical method used to determine the relationship between variables. "
        "Regression can predict continuous outcomes, such as estimating Delivery_Time_Hours based on Distance_km and Traffic_Level.\\n\\n"
        "C. Clustering: An unsupervised learning technique that groups similar data points. "
        "It can cluster geographical delivery zones or customer types with similar shipping behaviors to optimize regional hub placement.\\n\\n"
        "D. Optimization: Mathematical techniques used to find the best possible solution. "
        "In logistics, it applies to route optimization algorithms aimed at finding the shortest or cheapest path between multiple drop-off points.\\n\\n"
        "E. Correlation analysis: Evaluates the strength of relationship between two quantitative variables. "
        "This reveals insights such as whether an increase in Shipment_Weight_kg strongly correlates with higher Transportation_Cost.\\n\\n"
        "F. Descriptive statistics: Provides simple summaries about the sample and measures, such as mean and standard deviation. "
        "It is used to quickly calculate baseline KPIs like average delivery time."
    ).replace('\\n', '\n')
    add_paragraph(methodologies)
    
    doc.add_page_break()

    # 12. STRATEGIC ANALYTICAL ROADMAP
    doc.add_heading('12. STRATEGIC ANALYTICAL ROADMAP', level=1)
    roadmap = (
        "The project follows a structured end-to-end workflow:\\n\\n"
        "1. Data Collection: Sourcing and importing relevant public logistics datasets into a Python environment.\\n\\n"
        "2. Data Understanding: Reviewing dataset dimensions, data types, and understanding the context of each variable.\\n\\n"
        "3. Data Cleaning: Standardizing column names, fixing data formatting issues, and ensuring data consistency.\\n\\n"
        "4. Missing Value Handling: Identifying null records and imputing them with mean/median values or dropping them if negligible.\\n\\n"
        "5. Outlier Detection: Using boxplots and Z-scores to find and handle anomalous entries (e.g., unrealistic delivery times).\\n\\n"
        "6. Feature Engineering: Creating new useful columns, such as calculating 'Delay_Hours' from expected and actual times.\\n\\n"
        "7. Exploratory Data Analysis: Generating histograms, scatter plots, and bar charts to uncover initial patterns.\\n\\n"
        "8. KPI Calculation: Aggregating data to compute the established performance metrics.\\n\\n"
        "9. Correlation Analysis: Plotting heatmaps to identify strong relationships among numerical variables.\\n\\n"
        "10. Predictive Modeling: Setting up regression models to forecast delivery delays based on historical features.\\n\\n"
        "11. Route/Resource Optimization: Designing heuristic models to suggest optimal vehicle loading and routing.\\n\\n"
        "12. Business Insights: Translating statistical findings into plain-language observations.\\n\\n"
        "13. Recommendations: Proposing actionable strategies to logistics management based on the insights derived."
    ).replace('\\n', '\n')
    add_paragraph(roadmap)

    # 13. DATA CLEANING PLAN
    doc.add_heading('13. DATA CLEANING PLAN', level=1)
    cleaning = (
        "Ensuring data quality is a critical prerequisite for analysis. The future dataset will be cleaned using the following steps:\\n\\n"
        "• Missing values: Fields like 'Weather_Condition' may have blanks; these will be filled with the mode or a 'Unknown' category.\\n"
        "• Duplicate records: Exact duplicate rows resulting from system errors will be removed to prevent skewed counts.\\n"
        "• Incorrect data types: 'Order_Date' will be parsed into datetime objects, and numerical columns will be cast to floats or integers.\\n"
        "• Invalid values: Negative distances or costs will be flagged and removed or corrected.\\n"
        "• Outliers: Extreme values in 'Delivery_Time_Hours' will be capped or removed if they represent data entry errors.\\n"
        "• Inconsistent categorical values: String capitalization variations (e.g., 'van' vs 'Van') will be standardized.\\n"
        "• Date/time formatting: Standardizing timestamps to ISO format for uniform chronological analysis.\\n"
        "• Numerical normalization: Where appropriate for modeling, continuous variables like distance and cost may be scaled."
    ).replace('\\n', '\n')
    add_paragraph(cleaning)
    
    # 14. EXPLORATORY DATA ANALYSIS PLAN
    doc.add_heading('14. EXPLORATORY DATA ANALYSIS PLAN', level=1)
    eda = (
        "The EDA phase will utilize Python libraries including pandas, numpy, matplotlib, and seaborn to uncover operational realities:\\n\\n"
        "• Distribution of delivery times: Plotting histograms to see the typical spread of delivery duration.\\n"
        "• Shipment volume trends: Time-series line charts to observe weekly or monthly shipment peaks.\\n"
        "• Transportation cost analysis: Boxplots to compare cost distributions across different routes.\\n"
        "• Distance vs delivery time: Scatter plots to assess how linearly time increases with distance.\\n"
        "• Distance vs transportation cost: Analyzing if longer routes strictly dictate proportionally higher costs.\\n"
        "• Vehicle type comparison: Bar charts to evaluate average utilization and performance across different fleet categories.\\n"
        "• Delivery status analysis: Pie charts visualizing the proportion of 'On-Time' versus 'Late' shipments.\\n"
        "• Correlation matrix: A seaborn heatmap to summarize relationships between all numerical variables simultaneously."
    ).replace('\\n', '\n')
    add_paragraph(eda)

    doc.add_page_break()

    # 15. PYTHON CODE ILLUSTRATIONS
    doc.add_heading('15. PYTHON CODE ILLUSTRATIONS', level=1)
    add_paragraph("The following code snippets illustrate the foundational data operations using pandas, numpy, matplotlib, and seaborn.")
    
    code1 = """# A. Importing libraries
import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
import seaborn as sns

# B. Loading CSV data
df = pd.read_csv("logistics_data.csv")

# C. Checking missing values
print(df.isnull().sum())

# D. Removing duplicates
df.drop_duplicates(inplace=True)

# E. Basic descriptive statistics
print(df.describe())

# F. Calculating an on-time delivery KPI column (1 if on time, 0 if late)
df['Is_On_Time'] = np.where(df['Delivery_Time_Hours'] <= df['Expected_Delivery_Hours'], 1, 0)

# G. Creating a visualization (Histogram of Delivery Times)
plt.figure(figsize=(8, 5))
sns.histplot(df['Delivery_Time_Hours'], bins=30, kde=True)
plt.title("Distribution of Delivery Times")
plt.xlabel("Hours")
plt.ylabel("Frequency")
plt.show()

# H. Correlation analysis
plt.figure(figsize=(10, 6))
correlation_matrix = df[['Distance_km', 'Delivery_Time_Hours', 'Transportation_Cost', 'Shipment_Weight_kg']].corr()
sns.heatmap(correlation_matrix, annot=True, cmap='coolwarm')
plt.title("Correlation Matrix")
plt.show()"""
    add_code(code1)

    # 16. KPI CALCULATION EXAMPLE
    doc.add_heading('16. KPI CALCULATION EXAMPLE', level=1)
    code2 = """# Calculating Overall KPIs

# 1. On-Time Delivery Rate
on_time_rate = (df['Is_On_Time'].sum() / len(df)) * 100
print(f"On-Time Delivery Rate: {on_time_rate:.2f}%")

# 2. Average Delivery Time
avg_delivery_time = df['Delivery_Time_Hours'].mean()
print(f"Average Delivery Time: {avg_delivery_time:.2f} hours")

# 3. Average Transportation Cost
avg_transport_cost = df['Transportation_Cost'].mean()
print(f"Average Transportation Cost: ${avg_transport_cost:.2f}")

# 4. Vehicle Utilization Rate (assuming a hypothetical max capacity column exists)
df['Utilization_Rate'] = (df['Shipment_Volume'] / df['Max_Vehicle_Volume']) * 100
avg_utilization = df['Utilization_Rate'].mean()
print(f"Average Vehicle Utilization Rate: {avg_utilization:.2f}%")"""
    add_code(code2)
    add_paragraph("Explanation: The code utilizes pandas aggregation functions like .sum() and .mean() to efficiently compute the defined KPIs across the entire dataset.")

    doc.add_page_break()

    # 17. EXPECTED ANALYTICAL INSIGHTS
    doc.add_heading('17. EXPECTED ANALYTICAL INSIGHTS', level=1)
    insights = (
        "(Note: The following points represent POTENTIAL insights expected to be discovered during analysis, rather than actual findings.)\\n\\n"
        "• The analysis is expected to reveal which specific geographic routes consistently exhibit higher delivery times, pointing towards chronic traffic or infrastructure issues.\\n"
        "• It may highlight that external factors, such as specific Weather_Conditions, heavily influence delivery delays independent of distance.\\n"
        "• We expect to uncover which vehicle categories possess the highest and lowest utilization rates, suggesting areas for fleet optimization.\\n"
        "• The data is anticipated to show which intercity routes accumulate disproportionately higher transportation costs relative to shipment volume.\\n"
        "• Finally, correlation analysis is expected to clarify whether shipment distance linearly dictates delivery time and cost, or if other variables complicate this relationship."
    ).replace('\\n', '\n')
    add_paragraph(insights)

    # 18. EXPECTED OUTCOMES
    doc.add_heading('18. EXPECTED OUTCOMES', level=1)
    outcomes = (
        "Implementing this strategic analytical plan will yield significant benefits for logistics management:\\n\\n"
        "• Better visibility into day-to-day logistics performance.\\n"
        "• Accurate identification of operational bottlenecks causing delays.\\n"
        "• Improved resource allocation by matching vehicle capacities closely with shipment demands.\\n"
        "• Better route planning capabilities based on historical time and cost data.\\n"
        "• Reduced unnecessary transportation costs through identified efficiency gaps.\\n"
        "• Improved overall delivery reliability, thereby boosting customer satisfaction.\\n"
        "• A shift towards continuous, data-driven decision-making in supply chain operations."
    ).replace('\\n', '\n')
    add_paragraph(outcomes)

    # 19. FUTURE EXTENSION
    doc.add_heading('19. FUTURE EXTENSION', level=1)
    extension = (
        "In Weeks 2–4 of the internship, this foundational plan will be expanded upon. "
        "The project will progress into advanced data preprocessing and complex EDA utilizing dynamic visualizations. "
        "A major focus will be predictive delivery-time modeling, applying regression and tree-based machine learning models (such as Random Forest). "
        "These models will be rigorously evaluated using metrics like Mean Absolute Error (MAE), Root Mean Squared Error (RMSE), and R² score. "
        "\\n\\nFurthermore, prescriptive analytics will be explored to suggest algorithmic route optimization and dynamic resource allocation. "
        "The final deliverable could incorporate the development of an interactive dashboard to present these metrics and predictions to non-technical stakeholders."
    ).replace('\\n', '\n')
    add_paragraph(extension)

    # 20. LIMITATIONS
    doc.add_heading('20. LIMITATIONS', level=1)
    limitations = (
        "This strategic plan acknowledges several realistic limitations:\\n\\n"
        "• This is based on a hypothetical scenario; real-world logistics networks involve significantly more complexity.\\n"
        "• Dataset availability and data quality from public sources may not perfectly mirror proprietary enterprise data.\\n"
        "• Important missing operational variables (like driver fatigue or specific road closures) might limit analysis.\\n"
        "• Extreme traffic and weather variability are difficult to model purely from historical averages.\\n"
        "• Real-world route constraints (e.g., weight limits on bridges, delivery time-windows) may not be fully captured in open datasets.\\n"
        "• Predictive models rely on historical assumptions that may not hold true during unprecedented disruptions."
    ).replace('\\n', '\n')
    add_paragraph(limitations)

    doc.add_page_break()

    # 21. CONCLUSION
    doc.add_heading('21. CONCLUSION', level=1)
    conclusion = (
        "In conclusion, modern logistics operations face a multitude of challenges, from maintaining acceptable delivery times to controlling escalating transportation costs. "
        "This report outlines a comprehensive, data-driven strategy to address these issues by applying data science methodologies to operational shipment data. "
        "\\n\\nBy carefully defining and tracking Key Performance Indicators such as On-Time Delivery Rate and Average Transportation Cost, organizations can establish a quantifiable baseline of their network's efficiency. "
        "The proposed Python-based workflow—spanning data cleaning, exploratory data analysis, and correlation mapping—provides a structured mechanism for uncovering hidden operational inefficiencies. "
        "\\n\\nThe expected impact of executing this analytical roadmap is substantial. It moves logistics management away from instinct-based choices toward precise, evidence-backed resource allocation and route planning. "
        "While this plan is conceptual, the application of regression techniques and optimization models in subsequent phases holds the potential to significantly enhance delivery reliability. "
        "\\n\\nUltimately, integrating data science into supply chain management is no longer optional; it is a critical necessity for any logistics firm seeking to remain agile, cost-effective, and competitive in a demanding global market."
    ).replace('\\n', '\n')
    add_paragraph(conclusion)

    # 22. REFERENCES
    doc.add_heading('22. REFERENCES', level=1)
    refs = (
        "1. Kaggle. (n.d.). Supply Chain and Logistics Datasets. Retrieved from https://www.kaggle.com/\\n"
        "2. UCI Machine Learning Repository. (n.d.). Datasets related to delivery and demand forecasting. Retrieved from https://archive.ics.uci.edu/ml/index.php\\n"
        "3. The World Bank. (n.d.). Logistics Performance Index (LPI) Data. Retrieved from https://lpi.worldbank.org/\\n"
        "4. McKinney, W. (2010). Data Structures for Statistical Computing in Python (pandas documentation). Proceedings of the 9th Python in Science Conference.\\n"
        "5. Pedregosa, F., et al. (2011). Scikit-learn: Machine Learning in Python. Journal of Machine Learning Research, 12, 2825-2830."
    ).replace('\\n', '\n')
    add_paragraph(refs, justify=False)

    doc.save("c:/Users/Laxmi/OneDrive/Desktop/Internship/Laxmi_Ananda_Sanas_Week_1_Logistics_Data_Analyst_Report.docx")

if __name__ == '__main__':
    create_report()
