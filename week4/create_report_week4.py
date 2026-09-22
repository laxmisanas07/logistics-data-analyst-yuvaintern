import os
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE

def create_report():
    with open('stats.json', 'r') as f:
        stats = json.load(f)
        
    doc = Document()
    
    # Set page size to A4 and margins to 1 inch
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    def set_heading_style(level, size):
        try:
            style = doc.styles[f'Heading {level}']
        except KeyError:
            style = doc.styles.add_style(f'Heading {level}', WD_STYLE_TYPE.PARAGRAPH)
        font = style.font
        font.name = 'Calibri'
        font.size = Pt(size)
        font.bold = True
        font.color.rgb = RGBColor(0, 0, 0)
    
    set_heading_style(1, 16)
    set_heading_style(2, 13)
    
    title_style = doc.styles.add_style('CustomTitle', WD_STYLE_TYPE.PARAGRAPH)
    title_style.font.name = 'Calibri'
    title_style.font.size = Pt(24)
    title_style.font.bold = True
    title_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def add_paragraph(text, bold=False, justify=True, font_name='Calibri'):
        p = doc.add_paragraph()
        if justify:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run = p.add_run(text)
        run.bold = bold
        run.font.name = font_name
        return p

    def add_code(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(text)
        run.font.name = 'Consolas'
        run.font.size = Pt(10)
        return p

    # 1. COVER PAGE
    for _ in range(5):
        doc.add_paragraph()
    
    doc.add_paragraph("YuvaIntern\nLogistics Data Analyst Intern\nWeek 4 Task", style='CustomTitle')
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph("Predictive Modeling and Optimization in Logistics Using Python", style='CustomTitle')
    p.runs[0].font.size = Pt(20)
    
    for _ in range(10):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Student Name: Laxmi Ananda Sanas\nDate: September 2026").bold = True
    
    doc.add_page_break()
    
    # TOC
    doc.add_heading('Table of Contents', level=1)
    toc = (
        "1. COVER PAGE\\n2. EXECUTIVE SUMMARY\\n3. INTRODUCTION\\n4. PROBLEM DEFINITION\\n"
        "5. DATA PREPARATION FOR MODELING\\n6. PREDICTIVE MODELING METHODOLOGY\\n7. MODEL IMPLEMENTATION (PYTHON)\\n"
        "8. MODEL EVALUATION AND COMPARISON\\n9. HYPERPARAMETER TUNING\\n10. VISUALIZATION 1 — ACTUAL VS PREDICTED\\n"
        "11. VISUALIZATION 2 — RESIDUAL PLOT\\n12. VISUALIZATION 3 — MODEL COMPARISON\\n13. VISUALIZATION 4 — FEATURE IMPORTANCE\\n"
        "14. FEATURE IMPORTANCE ANALYSIS\\n15. LOGISTICS OPTIMIZATION STRATEGIES\\n16. BUSINESS INSIGHTS\\n"
        "17. LIMITATIONS\\n18. CONCLUSION\\n19. REFERENCES"
    ).replace('\\n', '\n')
    add_paragraph(toc, justify=False)
    
    doc.add_page_break()

    # 2. EXECUTIVE SUMMARY
    doc.add_heading('2. EXECUTIVE SUMMARY', level=1)
    summary = (
        "This report outlines the development of a predictive machine learning model to forecast logistics delivery times. "
        "Continuing from the data preprocessing (Week 2) and exploratory analysis (Week 3), this Week 4 project utilizes "
        "multiple regression techniques—Linear Regression, Decision Tree Regression, and Random Forest Regression—to predict "
        "the target variable based on shipment and operational features.\\n\\n"
        "Through rigorous data preparation involving one-hot categorical encoding and train-test splitting, models were trained "
        "and evaluated using Mean Absolute Error (MAE), Root Mean Squared Error (RMSE), and R-squared (R²). After initial testing, "
        "the Random Forest model emerged as the most robust predictor. Cross-validation and hyperparameter tuning were applied "
        "to optimize this model, resulting in an enhanced ability to accurately forecast transit times.\\n\\n"
        "Visualizations including actual-versus-predicted plots, residual tracking, and feature importance charts clearly highlight "
        "the model's performance and the driving factors behind delivery delays. Finally, practical logistics optimization strategies "
        "derived from these predictions are proposed, empowering data-driven decisions regarding resource allocation and route planning."
    ).replace('\\n', '\n')
    add_paragraph(summary)

    # 3. INTRODUCTION
    doc.add_heading('3. INTRODUCTION', level=1)
    intro = (
        "Predictive analytics transforms historical logistics data from a simple record of past events into a strategic tool for future planning. "
        "By anticipating how long a shipment will take based on specific parameters (like distance, vehicle type, and historical transit data), "
        "a logistics company can proactively manage its fleet, reduce idle time, and significantly enhance customer satisfaction. "
        "\\n\\nThis project demonstrates the end-to-end process of building such a predictive system using Python. It transitions from basic descriptive "
        "statistics into prescriptive optimization, showcasing the ultimate value of a fully matured data analytics pipeline."
    ).replace('\\n', '\n')
    add_paragraph(intro)

    # 4. PROBLEM DEFINITION
    doc.add_heading('4. PROBLEM DEFINITION', level=1)
    prob = (
        "The core logistics forecasting problem addressed here is the accurate prediction of delivery durations. Unpredictable delivery times "
        "lead to poor customer experiences and inefficient resource scheduling.\\n\\n"
        "• Target Variable: `Delivery_Time_Hours` (Continuous numerical variable).\\n"
        "• Predictor Variables (Features): `Distance_km`, `Vehicle_Type`, `Shipment_Weight_kg`, `Transportation_Cost`, `Origin`, and `Destination`.\\n\\n"
        "The objective is to train a model that minimizes the prediction error (MAE and RMSE) while explaining a high degree of the variance (R²)."
    ).replace('\\n', '\n')
    add_paragraph(prob)

    # 5. DATA PREPARATION FOR MODELING
    doc.add_heading('5. DATA PREPARATION FOR MODELING', level=1)
    prep = (
        "Before feeding the dataset into machine learning algorithms, specific transformations were required beyond the standard cleaning:\\n"
        "1. Feature Selection: Identifiers (`Shipment_ID`) and date variables (`Order_Date`) were dropped as they do not provide direct mathematical value to standard regression algorithms.\\n"
        "2. Categorical Encoding: Machine learning models require numerical input. Categorical columns (`Vehicle_Type`, `Origin`, etc.) were transformed using One-Hot Encoding (`pd.get_dummies`), which creates binary columns for each category.\\n"
        "3. Train-Test Splitting: To prevent overfitting and properly evaluate the model, the dataset was split—allocating 80% of the data for training the algorithms and 20% for testing them on unseen data."
    ).replace('\\n', '\n')
    add_paragraph(prep)
    
    # 6. PREDICTIVE MODELING METHODOLOGY
    doc.add_heading('6. PREDICTIVE MODELING METHODOLOGY', level=1)
    meth = (
        "Three distinct machine learning models were implemented to ensure a thorough comparative analysis:\\n"
        "• Linear Regression: Serves as a baseline model assuming a linear relationship between features and delivery time.\\n"
        "• Decision Tree Regression: A non-linear model that splits data based on feature thresholds, capable of capturing complex rules but prone to overfitting.\\n"
        "• Random Forest Regression: An ensemble learning method that constructs multiple decision trees and outputs the average prediction, effectively preventing overfitting and increasing accuracy."
    ).replace('\\n', '\n')
    add_paragraph(meth)

    # 7. MODEL IMPLEMENTATION (PYTHON)
    doc.add_heading('7. MODEL IMPLEMENTATION (PYTHON)', level=1)
    add_paragraph("The models were implemented using the `scikit-learn` library in Python:")
    add_code("""from sklearn.model_selection import train_test_split
from sklearn.linear_model import LinearRegression
from sklearn.tree import DecisionTreeRegressor
from sklearn.ensemble import RandomForestRegressor

# Split Data
X_train, X_test, y_train, y_test = train_test_split(X, y, test_size=0.2, random_state=42)

# Initialize Models
lr_model = LinearRegression()
dt_model = DecisionTreeRegressor(random_state=42)
rf_model = RandomForestRegressor(random_state=42)

# Train Models
lr_model.fit(X_train, y_train)
dt_model.fit(X_train, y_train)
rf_model.fit(X_train, y_train)""")

    # 8. MODEL EVALUATION AND COMPARISON
    doc.add_heading('8. MODEL EVALUATION AND COMPARISON', level=1)
    eval_text = "Models were evaluated on the test set using Mean Absolute Error (MAE), Root Mean Squared Error (RMSE), and R-squared (R²). The initial results before tuning were as follows:"
    add_paragraph(eval_text)
    
    init_res = stats['initial_results']
    table_eval = doc.add_table(rows=4, cols=4)
    table_eval.style = 'Table Grid'
    he = table_eval.rows[0].cells
    for i, t in enumerate(['Model', 'MAE (Hours)', 'RMSE (Hours)', 'R² Score']):
        he[i].text = t; he[i].paragraphs[0].runs[0].bold = True
        
    models = ['Linear Regression', 'Decision Tree', 'Random Forest']
    for i, m in enumerate(models):
        r = table_eval.rows[i+1].cells
        r[0].text = m
        r[1].text = f"{init_res[m]['MAE']:.2f}"
        r[2].text = f"{init_res[m]['RMSE']:.2f}"
        r[3].text = f"{init_res[m]['R2']:.2f}"
        
    add_paragraph("")
    add_paragraph("Analysis: The Random Forest model consistently outperformed the baseline Linear Regression and the single Decision Tree, exhibiting the lowest error margins and highest variance explanation.")
    
    doc.add_page_break()

    # 9. HYPERPARAMETER TUNING
    doc.add_heading('9. HYPERPARAMETER TUNING', level=1)
    tune = (
        "To further optimize the Random Forest model, hyperparameter tuning was conducted using `GridSearchCV`. "
        "This involved systematically testing combinations of parameters (such as `n_estimators`, `max_depth`, and `min_samples_split`) across cross-validated folds to find the optimal configuration.\\n\\n"
    ).replace('\\n', '\n')
    best_res = stats['best_rf_results']
    tune += (
        f"Optimized Random Forest Performance:\\n"
        f"• MAE: {best_res['MAE']:.2f} Hours\\n"
        f"• RMSE: {best_res['RMSE']:.2f} Hours\\n"
        f"• R² Score: {best_res['R2']:.2f}"
    ).replace('\\n', '\n')
    add_paragraph(tune)

    def embed_image(title, path, description, num):
        doc.add_heading(f"{num}. VISUALIZATION — {title.upper()}", level=1)
        if os.path.exists(path):
            doc.add_picture(path, width=Inches(6.0))
            add_paragraph(f"Figure {num-9}: {title}", justify=False)
        else:
            add_paragraph(f"[Image not found: {path}]", justify=False)
        add_paragraph(description)
        add_paragraph("")

    # Visualizations
    embed_image("Actual vs Predicted", "figures/actual_vs_predicted.png", 
        "This scatter plot maps the actual delivery times from the test set against the model's predicted times. Points lying on or near the dotted red line (perfect fit) indicate high prediction accuracy.", 10)

    embed_image("Residual Plot", "figures/residual_plot.png",
        "The residual plot shows the difference between the actual and predicted values. Ideally, these points should be randomly scattered around the horizontal zero line, indicating that the model's errors are random and normally distributed.", 11)

    embed_image("Model Comparison", "figures/model_comparison.png",
        "A dual-axis bar chart comparing the RMSE (error, lower is better) and R² (accuracy, higher is better) across the three tested models, visually confirming Random Forest as the superior algorithm.", 12)

    embed_image("Feature Importance", "figures/feature_importance.png",
        "This bar chart ranks the variables by their influence on the Random Forest model's predictions. It provides direct insight into which operational factors matter most when forecasting delivery times.", 13)

    # 14. FEATURE IMPORTANCE ANALYSIS
    doc.add_heading('14. FEATURE IMPORTANCE ANALYSIS', level=1)
    feat_text = "Based on the Random Forest model's internal calculations, the top influencing features are:\\n"
    for idx, f in enumerate(stats['top_features']):
        feat_text += f"{idx+1}. {f['Feature']} (Score: {f['Importance']:.3f})\\n"
    feat_text += "\\nThis analysis reveals that primary cost and distance metrics drastically overshadow categorical route identifiers in determining final delivery time."
    add_paragraph(feat_text.replace('\\n', '\n'))

    doc.add_page_break()

    # 15. LOGISTICS OPTIMIZATION STRATEGIES
    doc.add_heading('15. LOGISTICS OPTIMIZATION STRATEGIES', level=1)
    opt = (
        "Based on the predictive model, the following optimization strategies are proposed:\\n\\n"
        "1. Improved Resource Allocation: By predicting long transit times in advance, dispatchers can allocate faster or secondary vehicles to offset potential delays.\\n"
        "2. Dynamic Route Planning: If the model consistently flags specific Origin-Destination pairs as highly sensitive to delays, alternative routing can be pre-calculated.\\n"
        "3. Delivery Prioritization: The forecasted times can be used to tag 'at-risk' shipments, allowing warehouse staff to prioritize loading them earlier in the day.\\n"
        "4. Transportation Cost Management: By predicting transit times accurately, the company can avoid unnecessary expedited shipping fees previously used as a safety buffer."
    ).replace('\\n', '\n')
    add_paragraph(opt)

    # 16. BUSINESS INSIGHTS
    doc.add_heading('16. BUSINESS INSIGHTS', level=1)
    bi = (
        "The application of machine learning proves that delivery delays are largely predictable. "
        "The high R² achieved by the optimized Random Forest indicates that the current variables collected by the company "
        "are highly relevant. Operations should shift from a reactive 'track and trace' methodology to a proactive 'predict and prevent' paradigm."
    )
    add_paragraph(bi)

    # 17. LIMITATIONS
    doc.add_heading('17. LIMITATIONS', level=1)
    lim = (
        "• Model constraints: While Random Forest performs well, it cannot extrapolate beyond the distances and times seen in its training data.\\n"
        "• Missing Data: The model lacks real-time dynamic inputs (like live traffic or sudden weather events), which are often the true cause of unexpected delays.\\n"
        "• Simulated Sample: As this project utilizes a simulated dataset for demonstration, the specific MAE and R² values do not represent a real company's operational baseline."
    ).replace('\\n', '\n')
    add_paragraph(lim)

    # 18. CONCLUSION
    doc.add_heading('18. CONCLUSION', level=1)
    conc = (
        "The Week 4 logistics modeling project successfully bridged the gap between historical data analysis and forward-looking predictive intelligence. "
        "By applying a rigorous machine learning pipeline—from categorical encoding and train-test splitting to hyperparameter tuning and model evaluation—a robust "
        "Random Forest model was developed to forecast delivery times.\\n\\n"
        "The evaluation metrics clearly demonstrated the model's efficacy, significantly outperforming baseline linear approaches. More importantly, analyzing the model's "
        "feature importances provided actionable insights into the physical drivers of the supply chain. Visualizations such as the actual-versus-predicted and residual plots "
        "confirmed the statistical soundness of the approach.\\n\\n"
        "Ultimately, integrating predictive analytics into logistics operations enables powerful optimization strategies. By accurately forecasting transit times, "
        "organizations can optimize route planning, improve resource allocation, and drastically enhance the reliability of their deliveries."
    ).replace('\\n', '\n')
    add_paragraph(conc)

    # 19. REFERENCES
    doc.add_heading('19. REFERENCES', level=1)
    ref = (
        "1. Pedregosa, F., et al. (2011). Scikit-learn: Machine Learning in Python. Journal of Machine Learning Research, 12, 2825-2830.\\n"
        "2. Breiman, L. (2001). Random Forests. Machine Learning, 45(1), 5-32.\\n"
        "3. McKinney, W. (2010). Data Structures for Statistical Computing in Python (pandas).\\n"
        "4. Waskom, M. L. (2021). seaborn: statistical data visualization."
    ).replace('\\n', '\n')
    add_paragraph(ref, justify=False)

    doc.save("report/Laxmi_Ananda_Sanas_Week_4_Logistics_Modeling_Report.docx")

if __name__ == '__main__':
    create_report()
